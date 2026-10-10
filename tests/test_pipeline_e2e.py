"""Whole runs against the local fake BookMyShow / District / AllEvents."""

import hashlib
import threading
from datetime import datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import TODAY, browser_path, make_reference_workbook
from fakesites import FakeSite, FakeWeb
from s1scraper.config import Settings
from s1scraper.diagnose import run_diagnostics
from s1scraper.pipeline import Runner, build_notes
from s1scraper.sources import SOURCE_BY_KEY

THREE = ("bookmyshow", "district", "allevents")


def settings_for(tmp_path, use_browser=False, sources=THREE, **kw):
    wb = make_reference_workbook(tmp_path / "calendar.xlsx")
    s = Settings(workbook_path=str(wb), data_dir=str(tmp_path / "data"), speed="instant",
                 sources={k: k in sources for k in SOURCE_BY_KEY}, use_browser=use_browser,
                 window_preset="1 month", **kw)
    return s


def run(settings, web, **kw):
    return Runner(settings, rewrite=web.rewrite, unrewrite=web.unrewrite, today=TODAY, **kw).run()


@pytest.fixture(scope="module")
def plain_run(tmp_path_factory, fake_web):
    tmp = tmp_path_factory.mktemp("plain")
    s = settings_for(tmp)
    return s, run(s, fake_web)


def by_title(report):
    return {e.title: e for e in report.events}


def test_plain_run_finds_and_cleans_events(plain_run):
    s, rep = plain_run
    ev = by_title(rep)
    assert not rep.cancelled and not rep.error
    # BookMyShow: JSON state + cards + follow-on category page + plays
    for title in ("Aakash Gupta Live", "The Comedy Factory", "TOXIC – Abhishek Upmanyu Live", "Aadhe Adhure",
                  "Mumbai Food Truck Festival", "Rahul Dua – ALLOW ME", "Pune Jazz Night"):
        assert title in ev, title
    # District (slug fallback for Delhi) and AllEvents
    for title in ("Sunburn Arena ft. Alan Walker", "Sunday Gourmet Brunch", "Delhi Wine & Cheese Soirée",
                  "Kolkata Literary Meet 2026"):
        assert title in ev, title
    titles = set(ev)
    assert "Snow World" not in titles                      # /activities/ attraction
    assert "New Year's Eve Gala 2027" not in titles        # outside the window
    assert "Online Python Bootcamp" not in titles          # online-only
    assert not any(t.startswith("Chennai Sangeetham") for t in titles)   # JS-only page needs the browser
    assert rep.dropped.get("online-only") == 1


def test_details_are_merged_from_event_pages(plain_run):
    ev = by_title(plain_run[1])
    a = ev["Aakash Gupta Live"]
    assert a.start == datetime(2026, 10, 10, 20, 0) and a.organizer == "OML Entertainment"
    assert (a.price_min, a.price_max) == (999.0, 2499.0) and a.tier == "Premium" and a.activity_type == "Comedy"
    assert a.language == "Hindi, English" and a.duration == "1hr 30mins"
    assert a.platforms == ["BookMyShow", "AllEvents"]       # AllEvents copy merged in
    cf = ev["The Comedy Factory"]                           # embedded state only, separate time field
    assert cf.start == datetime(2026, 10, 14, 21, 0) and cf.organizer == "Comedy Factory Productions"
    assert cf.notes and "\n" not in cf.notes and len(cf.notes) <= 141
    toxic = ev["TOXIC – Abhishek Upmanyu Live"]
    assert toxic.organizer == "Source: BookMyShow"          # platform listed as organiser = unknown
    assert toxic.tier == "Luxury"
    rahul = ev["Rahul Dua – ALLOW ME"]                      # two-city tour page: Mumbai date chosen
    assert rahul.start == datetime(2026, 10, 28, 20, 0) and rahul.venue == "Shanmukhananda Hall"
    sunburn = ev["Sunburn Arena ft. Alan Walker"]           # React Server Component payload
    assert sunburn.organizer == "Percept Live" and (sunburn.price_min, sunburn.price_max) == (1999.0, 15000.0)
    assert sunburn.notes == "Alan Walker returns to India for a one-night-only arena show."   # one line
    assert a.notes == "Aakash Gupta is back with an all new hour of stand-up."
    assert ev["Delhi Wine & Cheese Soirée"].city == "Delhi NCR"
    assert ev["Delhi Wine & Cheese Soirée"].activity_type == "F&B"
    lit = ev["Kolkata Literary Meet 2026"]
    assert lit.is_free and lit.tier == "Standard" and lit.organizer == "Kolkata Lit Society"
    assert ev["Aadhe Adhure"].organizer == "Ank Theatre Group" and ev["Aadhe Adhure"].activity_type == "Theatre"


def test_refused_page_is_filled_from_another_platform(plain_run):
    pk = by_title(plain_run[1])["Prateek Kuhad – Silhouettes Tour"]   # BookMyShow answered 403
    assert pk.platforms == ["BookMyShow", "Zomato District"]
    assert pk.start == datetime(2026, 10, 20, 19, 0) and pk.organizer == "Only Much Louder"
    assert (pk.price_min, pk.price_max) == (1499.0, 6999.0)


def test_events_are_sorted_and_written(plain_run):
    s, rep = plain_run
    days = [max(e.start.date(), rep.window_start) for e in rep.events]
    assert days == sorted(days)                                  # by date (then city within a day)
    wb = load_workbook(rep.saved_to)
    rolling = wb["Rolling Calendar"]
    titles = [rolling.cell(row=r, column=2).value for r in range(4, rolling.max_row + 1)]
    assert len([t for t in titles if t]) == len(rep.events)
    assert rep.new_in_master == len(rep.events) - 1          # Aakash Gupta was typed into the Master by hand


def test_second_run_uses_cache_and_adds_nothing(plain_run, fake_web):
    s, first = plain_run
    hits_before = dict(fake_web.sites["in.bookmyshow.com"].hits)
    again = run(s, fake_web)
    assert again.new_in_master == 0 and len(again.events) == len(first.events)
    hits_after = fake_web.sites["in.bookmyshow.com"].hits
    detail = "/events/aakash-gupta-live/ET00410001"
    assert hits_after.get(detail) == hits_before.get(detail)   # event page served from the cache


@pytest.mark.skipif(browser_path() is None, reason="no Chromium/Chrome/Edge available")
def test_browser_reads_javascript_only_listings(tmp_path, fake_web, monkeypatch):
    if browser_path():
        monkeypatch.setenv("S1_BROWSER_PATH", browser_path())
    s = settings_for(tmp_path, use_browser=True, sources=("bookmyshow",), cities=["Chennai"])
    rep = run(s, fake_web)
    chennai = [e for e in rep.events if e.city == "Chennai"]
    assert len(chennai) == 12                          # all three pages of the infinite scroll
    assert all(e.organizer == "Madras Music Academy" and e.start.hour == 18 for e in chennai)
    assert rep.browser


def test_stop_leaves_the_workbook_untouched(tmp_path, fake_web):
    s = settings_for(tmp_path)
    before = hashlib.sha256(Path(s.workbook_path).read_bytes()).hexdigest()
    stop = threading.Event()

    def progress(frac, text):
        if frac > 0.1:
            stop.set()

    rep = run(s, fake_web, progress_fn=progress, stop_event=stop)
    assert rep.cancelled
    assert hashlib.sha256(Path(s.workbook_path).read_bytes()).hexdigest() == before


def test_a_refusing_site_is_skipped_and_reported(tmp_path, fake_web):
    angry = FakeSite("in.bookmyshow.com")
    angry.routes["/robots.txt"] = (200, "text/plain", "User-agent: *\nAllow: /\n")
    for slug in ("mumbai", "pune"):
        angry.routes[f"/explore/events-{slug}"] = (429, "text/html", "Too Many Requests")
    angry.start()
    web = FakeWeb({**fake_web.sites, "in.bookmyshow.com": angry})
    try:
        s = settings_for(tmp_path, sources=("bookmyshow", "district"))
        rep = run(s, web)
        assert not rep.cancelled
        assert any("stopped contacting" in w for w in rep.warnings)
        assert "Sunburn Arena ft. Alan Walker" in by_title(rep)     # other sources still delivered
        assert sum(angry.hits.values()) < 15                         # backed off quickly
    finally:
        angry.stop()


def test_diagnostics_report(tmp_path, fake_web):
    s = settings_for(tmp_path, sources=("bookmyshow", "district"))
    lines = []
    assert run_diagnostics(s, lines.append, rewrite=fake_web.rewrite, unrewrite=fake_web.unrewrite) == 0
    text = "\n".join(lines)
    assert "■ BookMyShow (Mumbai)" in text and "robots.txt allows the listing page: yes" in text
    assert "events found on listing:" in text and "fields found on event pages: date" in text
    reports = list((Path(s.data_dir) / "diagnostics").glob("*/report.txt"))
    assert reports and list(reports[0].parent.glob("bookmyshow/listing-*.html"))


def test_offline_run_leaves_the_workbook_alone(tmp_path):
    offline = FakeWeb({})                      # every host unreachable
    s = settings_for(tmp_path)
    before = hashlib.sha256(Path(s.workbook_path).read_bytes()).hexdigest()
    rep = run(s, offline)
    assert not rep.events and rep.failed_sources
    assert any("left unchanged" in w for w in rep.warnings) and not rep.saved_to
    assert hashlib.sha256(Path(s.workbook_path).read_bytes()).hexdigest() == before


def test_partial_failure_is_flagged_in_the_rolling_sheet(tmp_path, fake_web):
    web = FakeWeb({h: site for h, site in fake_web.sites.items() if h != "allevents.in"})
    rep = run(settings_for(tmp_path), web)
    assert rep.failed_sources == ["AllEvents"] and rep.events
    subtitle = load_workbook(rep.saved_to)["Rolling Calendar"]["A2"].value
    assert subtitle.startswith("⚠ Not read this run") and "AllEvents" in subtitle


def test_allevents_events_from_other_cities_are_dropped(tmp_path):
    from s1scraper.cities import resolve_cities
    from s1scraper.models import Event
    from s1scraper.pipeline import RunReport

    s = settings_for(tmp_path)
    evs = [Event(title=t, city="Mumbai", start=datetime(2026, 10, 11, 19), url=u, source="allevents",
                 platform="AllEvents", venue=v) for t, u, v in [
        ("Rangtarang 2026", "https://allevents.in/mehsana/rangtarang-2026-tickets/80009183438727",
         "Bliss Aqua World Resort, Mehsana"),
        ("Chittorgarh Bhajan Clubbing", "https://allevents.in/chittorgarh/chittorgarh-bhajan-clubbing/80003644859803",
         "Shreenath garden, RJ"),
        ("Moon Lamp", "https://allevents.in/mumbai/moon-lamp/3900030789365831", "Starbucks, Oshiwara"),
    ]]
    rep = RunReport(started=datetime(2026, 10, 7, 22), cities=["Mumbai"], sources=["AllEvents"])
    kept = Runner(s)._finalize(evs, resolve_cities(["Mumbai"]), datetime(2026, 10, 7).date(),
                               datetime(2026, 10, 14).date(), rep)
    assert [e.title for e in kept] == ["Moon Lamp"]


def test_online_overseas_and_placeholder_events_are_left_out():
    from s1scraper.models import Event
    from s1scraper.pipeline import is_abroad, is_online, is_placeholder_title

    assert is_online(Event(title="Founders Session", venue="Online event"))
    assert is_online(Event(title="Free Online English Classes", venue="Delhi"))
    assert is_online(Event(title="Career Webinar Recording", venue="Sales end soon"))
    assert not is_online(Event(title="VR Arena: Virtual Reality games", venue="Phoenix Mall"))
    assert is_abroad(Event(title="Sahaja Yoga Meditation", venue="Community Hall", address="Parramatta NSW, Australia"))
    assert is_abroad(Event(title="Leading SAFe 6.0 Training in Worcester, MA", url="https://www.eventbrite.com/e/x-1"))
    assert not is_abroad(Event(title="HD Navratri", url="https://www.eventbrite.com/e/hd-navratri-tickets-2000315676034"))
    # a Mumbai event on Eventbrite's UK site is still a Mumbai event
    assert not is_abroad(Event(title="Seed Business School Festival 2026 - Mumbai", venue="The St. Regis Mumbai",
                               url="https://www.eventbrite.co.uk/e/seed-business-school-festival-2026-mumbai-tickets-1"))
    assert [is_placeholder_title(t) for t in ("post", "Sdd", "Other Org", "October 11", "BTS", "Moon Lamp")] == [
        True, True, True, True, False, False]


def test_notes_fallback_reads_like_words():
    from s1scraper.models import Event

    ev = Event(title="Jashn-E-Qalam", categories=["Poetry", "age-bucket-8-to-15"], language="hindi|urdu")
    ev.activity_type = "Culture"
    assert build_notes(ev) == "Poetry in Hindi, Urdu."


# ------------------------------------------------- links, dates and register-by (real-run fixes)

def test_bookmyshow_links_point_at_the_right_section(plain_run):
    ev = by_title(plain_run[1])
    play = ev["Ek Tichi Goshta"]                      # built from the plays listing's data: /plays/, not /events/
    assert play.buy_link == "https://in.bookmyshow.com/plays/ek-tichi-goshta/ET00410011" and play.link_verified
    moved = ev["Hi Dosti Tutaychi Naay"]              # listed in the events grid; its page is under /plays/
    assert moved.buy_link == "https://in.bookmyshow.com/plays/hi-dosti-tutaychi-naay/ET00410012"
    assert moved.link_verified and moved.start.hour == 19
    assert ev["Aakash Gupta Live"].link_verified


def test_register_by_comes_from_the_site_and_never_after_the_start(plain_run):
    ev = by_title(plain_run[1])
    food = ev["Mumbai Food Truck Festival"]           # latest "sales end" of its two passes
    assert food.register_by == datetime.combine(TODAY + timedelta(days=13), time(23, 59))
    assert not food.register_by_inferred
    aakash = ev["Aakash Gupta Live"]                  # sales run past the start: capped at the start
    assert aakash.register_by == aakash.start and not aakash.register_by_inferred
    night = ev["Bollywood Night at Kitty Su"]         # "Registrations close on ..." in the page text
    assert (night.start.date() - night.register_by.date()).days == 1 and not night.register_by_inferred
    brunch = ev["Sunday Gourmet Brunch"]              # nothing published: open until it starts (shown in grey)
    assert brunch.register_by == brunch.start and brunch.register_by_inferred


def test_end_dates_are_trustworthy(plain_run):
    rep = plain_run[1]
    ev = by_title(rep)
    night = ev["Bollywood Night at Kitty Su"]         # 10 PM to 2 AM is one evening
    assert night.end is None and night.end_known
    season = ev["Kolkata Food Walks Season"]          # 5-month span: the last date is not printed
    assert not season.end_known
    assert "Weekly Kolkata Heritage Walk" not in ev    # began months ago, dates in the window unknown
    assert rep.dropped.get("repeats on dates not listed") == 1
    assert ev["TOXIC – Abhishek Upmanyu Live"].end.date() > ev["TOXIC – Abhishek Upmanyu Live"].start.date()

    ws = load_workbook(rep.saved_to)["Rolling Calendar"]
    head = [c.value for c in ws[3]]
    assert head.index("Register By") == head.index("Start Date") - 1
    rows = {r[1]: r for r in ws.iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)}
    assert rows["Kolkata Food Walks Season"][head.index("End Date")] == "Multiple dates"
    night_row = rows["Bollywood Night at Kitty Su"]
    assert night_row[head.index("End Date")] == night_row[head.index("Start Date")]
    brunch_cell = next(c for row in ws.iter_rows(min_row=4) for c in row
                       if row[1].value == "Sunday Gourmet Brunch" and c.column == head.index("Register By") + 1)
    assert brunch_cell.font.i and brunch_cell.number_format == "d mmm yyyy"     # inferred: grey italic


def test_listing_pages_beyond_the_first_and_no_needless_event_pages(plain_run, fake_web):
    rep = plain_run[1]
    ev = by_title(rep)
    jazz = ev["Kolkata Jazz Evening"]                 # only on page 2 of the listing
    assert jazz.organizer == "Calcutta Jazz Club" and jazz.price_min == 500
    assert fake_web.sites["allevents.in"].hits.get("/kolkata/kolkata-jazz-evening/80001235003", 0) == 0
    assert rep.detail_pages_skipped >= 1
    assert "Rangtarang Garba Nights 2026" not in ev    # AllEvents' Mumbai list, but held in Mehsana


def test_time_limit_keeps_listing_data(tmp_path, fake_web):
    s = settings_for(tmp_path, time_limit_minutes=1e-9)    # time is up before the first event page
    rep = run(s, fake_web)
    assert rep.detail_pages_over_time > 0 and rep.saved_to
    assert rep.listing_pages_cut > 0                       # only the first page of each list
    assert "Aakash Gupta Live" in by_title(rep)            # listing data still makes the sheet
    assert any("Time limit" in line for line in rep.summary_lines())


def test_time_budget_split():
    from s1scraper.pipeline import time_budget

    assert time_budget(0, 100.0) == (None, None)
    listing, detail = time_budget(60, 100.0)
    assert listing == 100.0 + 1800 and detail == 100.0 + 3600 - 120


def test_district_category_lists_are_read(plain_run):
    """The general District page shows only the next few days; concerts further out are on its category lists."""
    _, rep = plain_run
    ev = by_title(rep)
    rahman = ev["A.R. Rahman - The Wonderment Tour"]           # only on music-in-mumbai-book-tickets
    assert rahman.start.date() == TODAY + timedelta(days=24) and rahman.venue.startswith("Jio World Garden")
    assert "Dandiya Dhamaal Mumbai" in ev                     # a linked category list (navratri-in-...)
    assert "Weekend Slice Gig" not in ev                      # the this-weekend slice is not worth a request


def test_find_says_where_each_event_went(tmp_path, fake_web):
    lines = []
    s = settings_for(tmp_path)
    before = Path(s.workbook_path).read_bytes()
    Runner(s, rewrite=fake_web.rewrite, unrewrite=fake_web.unrewrite, today=TODAY, log_fn=lines.append,
           find=["A R Rahman", "Rangtarang Garba", "Imaginary Band"], write_excel=False).run()
    text = "\n".join(lines)
    assert 'Find "A R Rahman": in the sheet: "A.R. Rahman - The Wonderment Tour"' in text
    assert 'Find "Rangtarang Garba": not in the sheet.' in text
    assert "left out (other city): \"Rangtarang Garba Nights 2026\"" in text
    assert 'Find "Imaginary Band": not on any listing page this run read.' in text
    assert Path(s.workbook_path).read_bytes() == before          # a dry run leaves the workbook alone


def test_register_by_without_a_deadline_is_the_last_day_you_can_go():
    from s1scraper.models import Event
    from s1scraper.pipeline import settle_register_by

    run = Event(title="Clay Workshop", start=datetime(2026, 10, 9), end=datetime(2026, 11, 1))
    settle_register_by(run, TODAY)
    assert run.register_by == datetime(2026, 11, 1) and run.register_by_inferred     # not 9 Oct, already past
    sale = Event(title="Garba Nights", start=datetime(2026, 10, 9), end=datetime(2026, 10, 20),
                 register_by=datetime(2026, 10, 15))
    settle_register_by(sale, TODAY)
    assert sale.register_by == datetime(2026, 10, 15) and not sale.register_by_inferred   # published: kept
    many = Event(title="Open Mic", start=datetime(2026, 10, 3), end_known=False)
    settle_register_by(many, TODAY)
    assert many.register_by.date() == TODAY and many.register_by_inferred                 # not before today
