from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import make_reference_workbook
from s1scraper.classify import classify_activity, classify_tier
from s1scraper.config import Settings
from s1scraper.excel import HEADERS, INDEX_SHEET, MASTER_HEADERS, META_SHEET, is_locked, write_workbook
from s1scraper.models import Event
from s1scraper.pipeline import RunReport, build_notes

CITIES = ["Mumbai", "Pune", "Delhi NCR", "Bengaluru", "Kolkata", "Ahmedabad", "Chennai"]


def mk(title, city, start, platform="BookMyShow", lo=999.0, hi=None, cats=("Comedy Shows",), organizer="",
       code=None, **kw):
    url = kw.pop("url", f"https://in.bookmyshow.com/events/{title.lower().replace(' ', '-')}/{code or 'ET0' + str(abs(hash(title)) % 10**8)}")
    e = Event(title=title, city=city, start=start, platform=platform, platforms=[platform], url=url,
              links={platform: url}, price_min=lo, price_max=hi, categories=list(cats), organizer=organizer,
              source=platform.lower(), **kw)
    e.activity_type = classify_activity(e.categories, e.title)
    e.tier = classify_tier(e.price_min, e.price_max, e.is_free)
    e.organizer = e.organizer or f"Source: {platform}"
    e.notes = build_notes(e)
    return e


def report(started=datetime(2026, 10, 8, 15, 42), window=(date(2026, 10, 8), date(2026, 11, 8))):
    return RunReport(started=started, window_start=window[0], window_end=window[1], cities=CITIES,
                     sources=["BookMyShow", "District (Zomato)"],
                     per_source={"BookMyShow": {"discovered": 3, "kept": 3}})


@pytest.fixture
def settings(tmp_path):
    path = make_reference_workbook(tmp_path / "Rolling_Event_Calendar.xlsx")
    return Settings(workbook_path=str(path), data_dir=str(tmp_path / "data"), keep_backups=2)


EVENTS = lambda: [  # noqa: E731
    mk("Aakash Gupta Live", "Mumbai", datetime(2026, 10, 10, 20), code="ET00410001"),   # also typed by hand in Master
    mk("Prateek Kuhad Live", "Mumbai", datetime(2026, 10, 20, 19), lo=1499, hi=6999, cats=("Music Shows",),
       organizer="Only Much Louder", code="ET00410003"),
    mk("Pune Jazz Night", "Pune", datetime(2026, 10, 11, 21), lo=600, cats=("Music Shows",), code="ET00430001",
       venue="High Spirits Cafe", address="Koregaon Park, Pune"),
    mk("Delhi Wine Soirée", "Delhi NCR", datetime(2026, 11, 2, 19), platform="Zomato District", lo=4500,
       cats=("Food & Drinks",), url="https://www.district.in/events/delhi-wine-soiree-buy-tickets"),
]


def snapshot(ws, max_row, max_col):
    out = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            f, fl = cell.font, cell.fill
            out[(r, c)] = (cell.value, f.name, f.sz, f.b, str(f.color.rgb if f.color else None), fl.fill_type,
                           str(fl.fgColor.rgb), cell.number_format, cell.border.left.style,
                           cell.hyperlink.target if cell.hyperlink else None)
    return out


# StepOne's 13 columns -> the agreed 14: Register By goes in before Start Date (E), Ticket Platform(s) (K) goes
OLD_TO_NEW = {1: 1, 2: 2, 3: 3, 4: 4, 5: 6, 6: 7, 7: 8, 8: 9, 9: 10, 10: 11, 12: 12, 13: 13}


def test_existing_master_rows_keep_their_content(settings):
    path = Path(settings.workbook_path)
    m0 = load_workbook(path)["Master Calander "]
    dims = (m0.max_row, m0.max_column)
    before = snapshot(m0, *dims)
    merges = {str(r) for r in m0.merged_cells.ranges}

    res = write_workbook(settings, EVENTS(), report())
    assert any("removed the Ticket Platform(s) column" in n and "Register By" in n for n in res.notes)
    m1 = load_workbook(path)["Master Calander "]
    after = snapshot(m1, dims[0], 14)
    for (r, c), cell in before.items():
        if c in OLD_TO_NEW:
            assert after[(r, OLD_TO_NEW[c])] == cell, (r, c)    # every other cell keeps its value and look
    assert [m1.cell(row=3, column=c).value for c in range(1, 15)] == MASTER_HEADERS
    assert merges <= {str(r) for r in m1.merged_cells.ranges}  # title and month bands still span A:M
    assert not [c.coordinate for row in m1.iter_rows() for c in row if c.hyperlink]
    assert m1.cell(row=6, column=5).fill.fgColor.rgb == m1.cell(row=6, column=6).fill.fgColor.rgb   # row look

    snap = snapshot(m1, m1.max_row, 14)
    res = write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15, 9)))
    assert not res.notes                                       # nothing left to change: append-only from here
    assert snapshot(load_workbook(path)["Master Calander "], m1.max_row, 14) == snap


def test_master_from_an_older_scraper_version_is_tidied(settings):
    """A Master that already got the old Link / Added On columns (and links in Ticket Platform(s))."""
    path = Path(settings.workbook_path)
    wb = load_workbook(path)
    m = wb["Master Calander "]
    m["N3"], m["O3"] = "Link", "Added On"
    m["B12"], m["E12"], m["K12"] = "Old Scraper Row", datetime(2026, 10, 12), "BookMyShow"
    m["N12"].hyperlink = "https://in.bookmyshow.com/events/old-scraper-row/ET00999999"
    m["K12"].hyperlink = "https://in.bookmyshow.com/events/old-scraper-row/ET00999999"
    m["O12"] = datetime(2026, 10, 1)
    m["A11"] = "▌ SCRAPER UPDATE · 01 OCT 2026 · 1 NEW EVENT"
    m.merge_cells("A11:O11")
    m.column_dimensions["N"].width, m.column_dimensions["O"].width = 38, 11
    m.freeze_panes = "C4"
    wb.save(path)

    write_workbook(settings, EVENTS(), report())
    m = load_workbook(path)["Master Calander "]
    assert [m.cell(row=3, column=c).value for c in range(1, 15)] == MASTER_HEADERS
    assert m.cell(row=3, column=15).value is None
    assert "A11:N11" in {str(r) for r in m.merged_cells.ranges}          # the old band spans the new width
    assert m["B12"].value == "Old Scraper Row" and m["F12"].value == datetime(2026, 10, 12)
    assert m["N12"].value == datetime(2026, 10, 1)                       # Added On kept, one column left
    assert m.column_dimensions["N"].width == 11 and m.freeze_panes == "C4"
    assert not [c.coordinate for row in m.iter_rows() for c in row if c.hyperlink]


def test_client_sheet_has_no_prices_platforms_or_links(settings):
    write_workbook(settings, EVENTS(), report())
    ws = load_workbook(settings.workbook_path)["Rolling Calendar"]
    header = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
    assert header == HEADERS and len(HEADERS) == 12
    assert not {"Price Range", "Organizer", "Ticket Platform(s)", "Link"} & set(header)
    assert header[4:6] == ["Register By", "Start Date"]
    assert not [c.coordinate for row in ws.iter_rows() for c in row if c.hyperlink]
    texts = [str(c.value) for row in ws.iter_rows(min_row=4) for c in row if c.value]
    assert not any("bookmyshow.com" in t or "₹" in t for t in texts)


def test_new_events_are_appended_once(settings):
    path = Path(settings.workbook_path)
    r1 = write_workbook(settings, EVENTS(), report())
    assert r1.new_in_master == 3          # Aakash Gupta was already typed into the Master by hand
    r2 = write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15, 9)))
    assert r2.new_in_master == 0

    m = load_workbook(path)["Master Calander "]
    assert [m.cell(row=3, column=c).value for c in range(1, 15)] == MASTER_HEADERS
    assert m.cell(row=10, column=1).value.startswith("▌ SCRAPER UPDATE · 08 OCT 2026 · 3 NEW EVENTS")
    assert "A10:N10" in {str(r) for r in m.merged_cells.ranges}
    bands = [m.cell(row=r, column=1).value for r in (11, 13, 15)]
    assert bands == ["MUMBAI · 1 EVENT", "PUNE · 1 EVENT", "DELHI NCR · 1 EVENT"]   # city by city
    rows = [[m.cell(row=r, column=c).value for c in range(1, 15)] for r in (12, 14, 16)]
    assert [r[0] for r in rows] == [4, 5, 6]                     # S.No. continues after the hand-typed 3
    assert [r[1] for r in rows] == ["Prateek Kuhad Live", "Pune Jazz Night", "Delhi Wine Soirée"]
    pune = rows[1]
    assert pune[4] == pune[5] == datetime(2026, 10, 11)          # no deadline published: Register By = start
    assert m.cell(row=14, column=6).number_format == "d mmm yyyy"
    assert pune[8] == "High Spirits Cafe, Koregaon Park"
    assert pune[9] == "Source: BookMyShow"                       # organiser fallback names the source
    assert rows[0][10] == "₹1,499 – ₹6,999" and rows[0][3] == "Premium"   # the price stays: it sets the tier
    assert pune[13] == datetime(2026, 10, 8)
    assert not [c.coordinate for row in m.iter_rows() for c in row if c.hyperlink]
    assert m.max_row == 16


def test_rows_deleted_by_hand_are_not_added_again(settings):
    path = Path(settings.workbook_path)
    write_workbook(settings, EVENTS(), report())
    wb = load_workbook(path)
    wb["Master Calander "].delete_rows(12)          # StepOne removes "Prateek Kuhad Live"
    wb.save(path)
    assert write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15))).new_in_master == 0


def test_rolling_sheet_is_rebuilt_each_run(settings):
    path = Path(settings.workbook_path)
    write_workbook(settings, EVENTS(), report())
    write_workbook(settings, EVENTS()[1:2], report(started=datetime(2026, 10, 15)))   # only one event now
    wb = load_workbook(path)
    ws = wb["Rolling Calendar"]
    assert [ws.cell(row=3, column=c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
    titles = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=2).value]
    assert titles == ["Prateek Kuhad Live"]
    assert ws.freeze_panes == "C4" and ws.auto_filter.ref.startswith("A3:L")
    assert ws["A4"].value == "▌ MUMBAI · 1 EVENT"
    assert wb.active.title == "Rolling Calendar"


def test_rolling_sheet_goes_city_by_city_in_date_order(settings):
    events = EVENTS() + [mk("Mumbai Early Show", "Mumbai", datetime(2026, 10, 9, 18), code="ET00410099")]
    write_workbook(settings, events, report())
    ws = load_workbook(settings.workbook_path)["Rolling Calendar"]
    lines = [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value) for r in range(4, ws.max_row + 1)]
    assert lines == [
        ("▌ MUMBAI · 3 EVENTS", None), (1, "Mumbai Early Show"), (2, "Aakash Gupta Live"), (3, "Prateek Kuhad Live"),
        ("▌ PUNE · 1 EVENT", None), (4, "Pune Jazz Night"),
        ("▌ DELHI NCR · 1 EVENT", None), (5, "Delhi Wine Soirée"),
    ]


def test_default_workbook_is_just_master_and_rolling(settings):
    write_workbook(settings, EVENTS(), report())
    write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15)))
    wb = load_workbook(settings.workbook_path)
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible == ["Guide", "Master Calander ", "Rolling Calendar"]
    assert wb["Run Log"].sheet_state == "hidden" and wb["Run Log"].max_row == 3


def test_city_tabs_summary_and_run_log(settings):
    path = Path(settings.workbook_path)
    settings.city_tabs = settings.summary_tab = True
    write_workbook(settings, EVENTS(), report())
    write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15)))
    wb = load_workbook(path)
    names = wb.sheetnames
    assert names[:4] == ["Guide", "Master Calander ", "Rolling Calendar", "Summary"]
    assert names[4:11] == CITIES and names[11] == "Run Log" and wb["Run Log"].sheet_state == "hidden"
    assert wb[INDEX_SHEET].sheet_state == "hidden" and wb[META_SHEET].sheet_state == "hidden"
    mumbai = [wb["Mumbai"].cell(row=r, column=2).value for r in range(4, wb["Mumbai"].max_row + 1)]
    assert [t for t in mumbai if t] == ["Aakash Gupta Live", "Prateek Kuhad Live"]
    assert "No events found in Kolkata" in wb["Kolkata"]["A4"].value
    summary = wb["Summary"]
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("COUNTIFS('Rolling Calendar'!$H$4:$H$" in f for f in formulas)    # City / Cities column
    log = wb["Run Log"]
    assert log.max_row == 3 and log.cell(row=3, column=4).value == 4


def test_sheet_names_the_user_owns_are_never_replaced(settings):
    path = Path(settings.workbook_path)
    settings.city_tabs = True
    wb = load_workbook(path)
    wb.create_sheet("Mumbai")["A1"] = "StepOne's own notes"
    wb.save(path)
    write_workbook(settings, EVENTS(), report())
    write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15)))
    wb = load_workbook(path)
    assert wb["Mumbai"]["A1"].value == "StepOne's own notes"
    assert "Mumbai (Rolling)" in wb.sheetnames


def test_turning_off_city_tabs_removes_only_generated_tabs(settings):
    settings.city_tabs = settings.summary_tab = True
    write_workbook(settings, EVENTS(), report())
    settings.city_tabs = settings.summary_tab = False
    write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, 15)))
    wb = load_workbook(settings.workbook_path)
    assert not set(CITIES) & set(wb.sheetnames) and "Summary" not in wb.sheetnames
    assert "Master Calander " in wb.sheetnames and "Guide" in wb.sheetnames


def test_new_workbook_is_created_when_missing(tmp_path):
    s = Settings(workbook_path=str(tmp_path / "new" / "calendar.xlsx"), data_dir=str(tmp_path / "data"))
    res = write_workbook(s, EVENTS(), report())
    assert res.new_in_master == 4 and any("Created a new workbook" in w for w in res.warnings)
    wb = load_workbook(res.saved_to)
    assert wb.sheetnames[:3] == ["Guide", "Master Calander ", "Rolling Calendar"]
    master = wb["Master Calander "]
    assert [master.cell(row=3, column=c).value for c in range(1, 16)] == MASTER_HEADERS + [None]
    assert not res.notes                                          # a new Master already has the agreed columns


def test_locked_workbook_is_saved_next_to_it(settings):
    path = Path(settings.workbook_path)
    lock = path.with_name("~$" + path.name)       # what Excel creates while the file is open
    lock.write_text("x")
    try:
        assert is_locked(path)
        res = write_workbook(settings, EVENTS(), report())
        assert res.saved_to != path and res.saved_to.exists() and "(updated " in res.saved_to.name
        assert any("open in another program" in w for w in res.warnings)
        assert "Rolling Calendar" not in load_workbook(path).sheetnames     # original untouched
    finally:
        lock.unlink()


def test_backups_are_kept_and_pruned(settings):
    for day in (8, 9, 10, 11):
        write_workbook(settings, EVENTS(), report(started=datetime(2026, 10, day)))
    backups = list((Path(settings.data_dir) / "backups").glob("*.xlsx"))
    assert 1 <= len(backups) <= settings.keep_backups
