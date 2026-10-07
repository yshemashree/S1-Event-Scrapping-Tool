from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from fakesites import start_fake_web  # noqa: E402

TODAY = date(2026, 10, 8)


@pytest.fixture(scope="session")
def fake_web():
    web = start_fake_web(TODAY)
    yield web
    web.stop()


def browser_path():
    """A Chromium the tests may drive, or None (browser tests are then skipped)."""
    import os

    candidate = os.environ.get("S1_BROWSER_PATH", "")
    if candidate and Path(candidate).exists():
        return candidate
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None
    with sync_playwright() as p:
        for channel in ("msedge", "chrome", None):   # same order the app tries
            try:
                b = p.chromium.launch(channel=channel) if channel else p.chromium.launch()
                b.close()
                return ""   # the app's own browser discovery will find it
            except Exception:  # noqa: BLE001
                continue
    return None


def make_reference_workbook(path: Path) -> Path:
    """A workbook shaped like StepOne's: Guide + 'Master Calander ' with title, subtitle,
    header, merged month bands, hyperlinks and hand-typed rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    g = wb.active
    g.title = "Guide"
    g["A1"] = "COLOUR KEY & REGION GUIDE"
    g["A1"].font = Font(name="Arial", size=12, bold=True, color="FFF9F6F0")
    g["A1"].fill = PatternFill("solid", start_color="FF1A1A1A")
    g.merge_cells("A1:D1")
    m = wb.create_sheet("Master Calander ")
    m["A1"] = "ROLLING LUXURY & ULTRA PREMIUM EVENTS CALENDAR · 11 JUNE 2026 – FEBRUARY 2027"
    m.merge_cells("A1:M1")
    m["A2"] = "Sources: BookMyShow · Zomato District · Last Updated: Jun 2026"
    m.merge_cells("A2:M2")
    headers = ["S.No.", "Event Name", "Activity Type", "Tier", "Start Date", "End Date", "City / Cities", "Venue",
               "Organizer", "Price Range", "Ticket Platform(s)", "Notes", "Region"]
    thin = Side(style="thin")
    for i, h in enumerate(headers, start=1):
        c = m.cell(row=3, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFF9F6F0")
        c.fill = PatternFill("solid", start_color="FF3A3A3A")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(horizontal="center")
    m["A4"] = "▌ OCTOBER 2026"
    m.merge_cells("A4:M4")
    rows = [
        [1, "Women's Indian Open Golf 2026", "Recreational Sports", "Luxury", datetime(2026, 10, 1),
         datetime(2026, 10, 1), "Gurugram", "DLF Golf & Country Club, Gurugram", "PGTI",
         "Free spectator – ₹1,500 (hospitality)", "DreamSetGo", "Women's golf tour event.", "India"],
        [2, "Aakash Gupta Live", "Comedy", "Premium", datetime(2026, 10, 10), datetime(2026, 10, 10), "Mumbai",
         "St. Andrew's Auditorium", "BookMyShow", "₹999 – ₹2,499", "BookMyShow", "Typed in by hand.", "India"],
        [3, "Sonny Fodera | Playa Pacha – Dubai", "Music", "Luxury", datetime(2026, 10, 10), datetime(2026, 10, 10),
         "Dubai, UAE", "FIVE Luxe JBR", "Pacha / FIVE Hotels", "AED 200–700", "Ticketmaster.ae", "Beach club.", "Intl"],
    ]
    for r, values in enumerate(rows, start=5):
        for c, v in enumerate(values, start=1):
            cell = m.cell(row=r, column=c, value=v)
            cell.font = Font(name="Arial", size=9, color="FF3A3A3A")
            cell.fill = PatternFill("solid", start_color="FFFAF0CC")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if isinstance(v, datetime):
                cell.number_format = "d mmm yyyy"
    m["K7"].hyperlink = "http://ticketmaster.ae"
    m["A8"] = "▌ NOVEMBER 2026"
    m.merge_cells("A8:M8")
    m["B9"] = "Rann Utsav"
    m["E9"] = "5th Nov"
    wb.save(path)
    return path
